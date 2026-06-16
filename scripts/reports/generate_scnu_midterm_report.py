from __future__ import annotations

# pyright: reportMissingImports=false

from dataclasses import dataclass
from pathlib import Path
from typing import Final

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT: Final = Path(__file__).resolve().parents[2]
OUT_DIR: Final = ROOT / "docs" / "reports"
OUT_PATH: Final = OUT_DIR / "2026_SCNU_창업동아리_중간점검_보고서.docx"
FONT: Final = "Malgun Gothic"
BLACK: Final = "000000"
HEADER_FILL: Final = "F2F4F7"
ACCENT_FILL: Final = "E8EEF5"


@dataclass(frozen=True, slots=True)
class Activity:
    date: str
    result: str
    detail: str


@dataclass(frozen=True, slots=True)
class BudgetLine:
    category: str
    allocated: str
    spent: str
    balance: str


ACTIVITIES: Final = (
    Activity("4.17.", "창업동아리 오리엔테이션", "운영 일정, 지원금 집행 기준, 중간점검 방향 확인"),
    Activity("5.20.~22.", "첨단소재캠프", "창업 아이디어 발굴 및 창업 프로그램 참여"),
    Activity("5.28.~29.", "리벨업캠프", "창업동아리 발표, 컨설팅, 아이템 고도화 방향 검토"),
    Activity("6.03.~04.", "1차 시제품 설계 확정", "코치 예약, 식단 코칭, 운동 의류 추천 MVP 흐름 확정"),
    Activity("6.04.", "MVP 시제품 구현", "한국어 모바일 앱, 운영자 웹, 백엔드 API, Docker 로컬 배포 구현"),
    Activity("6.04.", "배포·지식재산권 서류 정리", "개인정보·약관·보안·저작권 등록 증빙 패키지 초안 작성"),
    Activity("6.16.", "중간점검 자료 작성", "GitHub 업로드 준비, 중간보고서, 향후 추진 계획 정리"),
)

BUDGET: Final = (
    BudgetLine("재료비", "300,000", "0", "300,000"),
    BudgetLine("연구개발비", "500,000", "0", "500,000"),
    BudgetLine("일반운영비", "200,000", "160,000", "40,000"),
    BudgetLine("합계", "1,000,000", "160,000", "840,000"),
)

PLAN_TEXT: Final = (
    "시제품 개발 계획\n"
    "Supabase 실데이터 연결, 핵심 플로우 검증, TestFlight/내부 테스트 배포 준비.\n\n"
    "활동 목표\n"
    "7월 중 파일럿 5~10명 피드백 수집 후 시장검증 단계로 전환.\n\n"
    "산출물\n"
    "한국어 앱·운영자 웹·API/Docker 패키지·시연 영상·보안/저작권 증빙.\n\n"
    "지원금 사용 계획\n"
    "재료비=화면/테스트 자료, 연구개발비=DB/배포/AI 검증, 운영비=멘토링/발표/영상 제작.\n\n"
    "요청사항\n"
    "도메인·서버·내부 테스트·홍보물 예산 조정 또는 추가 지원 검토 요청."
)

RISK_TEXT: Final = (
    "실서비스 검증과 앱스토어 배포는 미완료이나 한국어 MVP, 운영자 웹, API, Docker 로컬 배포, "
    "개인정보·보안·저작권 증빙 문서는 구축했다. 남은 기간 Supabase 실데이터, 파일럿 테스트, 스토어 테스트 배포를 진행한다."
)


def _set_run_font(run, *, size: float, bold: bool = False, color: str = BLACK) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _set_cell_margins(cell, *, top: int = 35, bottom: int = 35, start: int = 90, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _dxa(inches: float) -> int:
    return round(inches * 1440)


def _cell_text(cell, text: str, *, bold: bool = False, size: float = 8.4, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold)


def _set_table_geometry(table, widths: tuple[float, ...]) -> None:
    width_values = tuple(_dxa(width) for width in widths)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(width_values)))

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width in width_values:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, width in enumerate(width_values):
            cell = row.cells[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def _paragraph(doc, text: str = "", *, size: float = 9.0, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.02
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, bold=bold)
    return paragraph


def _add_title(doc) -> None:
    paragraph = _paragraph(doc, "2026년 SCNU 창업동아리 중간점검 보고서", size=15.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.paragraph_format.space_after = Pt(4)


def _add_basic_info(doc) -> None:
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    _set_table_geometry(table, (1.0, 2.4, 1.0, 2.1))
    rows = (
        ("동아리명", "Fitconnect", "대표자명", "최성원"),
        ("아이템명", "Fitconnect", "현재단계", "☐ 아이디어   ☑ 시제품 제작\n☐ 시장검증   ☐ 매출 발생"),
    )
    for row_index, values in enumerate(rows):
        for col_index, text in enumerate(values):
            cell = table.cell(row_index, col_index)
            label_cell = col_index in (0, 2)
            _cell_text(cell, text, bold=label_cell, align=WD_ALIGN_PARAGRAPH.CENTER)
            if label_cell:
                _shade(cell, HEADER_FILL)


def _add_progress(doc) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    _set_table_geometry(table, (1.1, 0.9, 4.5))
    _cell_text(table.cell(0, 0), "목표 대비\n수행률", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(table.cell(0, 0), HEADER_FILL)
    _cell_text(table.cell(0, 1), "70%", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(table.cell(0, 1), ACCENT_FILL)
    _cell_text(
        table.cell(0, 2),
        "한국어 모바일 MVP, 운영자 웹, 백엔드 API, Docker 로컬 배포와 제출 문서를 구축했다. "
        "Supabase 실데이터 전환, 파일럿 사용자 검증, 스토어 배포는 진행 예정이다.",
        size=9,
    )


def _add_activity_table(doc) -> None:
    _paragraph(doc, "주요 활동 내용", size=10, bold=True)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    _set_table_geometry(table, (0.95, 1.7, 3.85))
    for index, title in enumerate(("일자", "성과", "활동 내용")):
        _cell_text(table.cell(0, index), title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(table.cell(0, index), HEADER_FILL)
    for activity in ACTIVITIES:
        row = table.add_row()
        _cell_text(row.cells[0], activity.date, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(row.cells[1], activity.result, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(row.cells[2], activity.detail)


def _add_budget_table(doc) -> None:
    _paragraph(doc, "지원금 집행현황(6월말기준)", size=10, bold=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    _set_table_geometry(table, (1.4, 1.7, 1.7, 1.7))
    for index, title in enumerate(("집행비목", "배정지원금(원)", "집행금액(원)", "잔액")):
        _cell_text(table.cell(0, index), title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(table.cell(0, index), HEADER_FILL)
    for item in BUDGET:
        row = table.add_row()
        bold = item.category == "합계"
        for index, text in enumerate((item.category, item.allocated, item.spent, item.balance)):
            _cell_text(row.cells[index], text, bold=bold, align=WD_ALIGN_PARAGRAPH.CENTER)
        if bold:
            for cell in row.cells:
                _shade(cell, ACCENT_FILL)


def _add_text_section(doc, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _set_table_geometry(table, (1.1, 5.4))
    _cell_text(table.cell(0, 0), label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _shade(table.cell(0, 0), HEADER_FILL)
    _cell_text(table.cell(0, 1), text, size=7.8)


def _add_signature(doc) -> None:
    _paragraph(doc, "우리 창업동아리는 위 내용과 같이 중간보고서를 제출합니다.", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _paragraph(doc, "2026년  06월  17일", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _paragraph(doc, "동아리 대표 :        최성원        (서명/날인)", size=9, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _paragraph(doc, "국립순천대학교 창업지원단장 귀하", size=12, bold=True)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.25)
    section.bottom_margin = Inches(0.2)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    styles = doc.styles["Normal"]
    styles.font.name = FONT
    styles.font.size = Pt(8.4)
    styles.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _add_title(doc)
    _add_basic_info(doc)
    _add_progress(doc)
    _add_activity_table(doc)
    _add_budget_table(doc)
    _add_text_section(doc, "향후\n활동 계획", PLAN_TEXT)
    _add_text_section(doc, "목표\n미달사유", RISK_TEXT)
    _add_signature(doc)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
